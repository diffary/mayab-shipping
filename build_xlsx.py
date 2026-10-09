"""Build stuck_parcels.xlsx from source.xlsx: adds Rules, helper columns, Summary.
Upload the result to Google Drive and open as Google Sheets."""
import openpyxl
from openpyxl.styles import Font, PatternFill
from datetime import datetime

wb = openpyxl.load_workbook("source.xlsx")
data = wb.worksheets[0]
data.title = "Data"
N = data.max_row  # 1001

# ---- Rules sheet -------------------------------------------------------
rules = wb.create_sheet("Rules", 0)
rules["A1"], rules["B1"] = "today", datetime(2026, 10, 15)
rules["A2"], rules["B2"] = "idle_days", 7
rules["A3"], rules["B3"] = "customs_days", None
rules["A5"] = "Edit the yellow cells. Leave customs_days empty to switch the customs rule off."
rules["B1"].number_format = "dd.mm.yyyy"
for c in ("B1", "B2", "B3"):
    rules[c].fill = PatternFill("solid", fgColor="FFFF00")
    rules[c].font = Font(color="0000FF")
rules.column_dimensions["A"].width = 16

# ---- helper columns in Data -------------------------------------------
data["H1"], data["I1"], data["J1"] = "days idle", "reason", "priority"
for row in range(2, N + 1):
    r = row
    data[f"H{r}"] = f'=IF(G{r}="","",Rules!$B$1-G{r})'
    data[f"I{r}"] = (
        f'=IF(F{r}="Address issue","address issue",'
        f'IF(OR(F{r}="Delivered",F{r}="Returned to sender",G{r}=""),"",'
        f'IF(AND(Rules!$B$3<>"",F{r}="Held at customs",H{r}>=Rules!$B$3),"customs >= "&Rules!$B$3&" days",'
        f'IF(H{r}>=Rules!$B$2,"idle >= "&Rules!$B$2&" days",""))))'
    )
    # rank: address issues first, then longest idle; ties share a number
    data[f"J{r}"] = (
        f'=IF(I{r}="","",'
        f'IF(F{r}="Address issue",0,COUNTIF($I$2:$I${N},"address issue"))'
        f'+COUNTIFS($I$2:$I${N},"?*",$F$2:$F${N},IF(F{r}="Address issue","Address issue","<>Address issue"),$H$2:$H${N},">"&H{r})'
        f'+1)'
    )
for c in ("H1", "I1", "J1"):
    data[c].font = Font(bold=True)

# ---- Summary sheet -----------------------------------------------------
s = wb.create_sheet("Summary", 1)
s["A1"], s["B1"], s["C1"], s["D1"] = "Carrier", "Stuck", "Total", "Share"
for i, carrier in enumerate(["AndesCargo", "SwiftParcel", "PostLink"], start=2):
    s[f"A{i}"] = carrier
    s[f"B{i}"] = f'=COUNTIFS(Data!$C$2:$C${N},A{i},Data!$I$2:$I${N},"?*")'
    s[f"C{i}"] = f'=COUNTIF(Data!$C$2:$C${N},A{i})'
    s[f"D{i}"] = f'=IF(C{i}=0,0,B{i}/C{i})'
    s[f"D{i}"].number_format = "0%"
s["A5"], s["B5"], s["C5"] = "Total", "=SUM(B2:B4)", "=SUM(C2:C4)"
s["A7"] = "today:"
s["B7"] = "=Rules!B1"
s["B7"].number_format = "dd.mm.yyyy"
s["A8"] = "Stuck list: Data sheet, filter column I not empty, sort by column J."
for c in ("A1", "B1", "C1", "D1", "A5"):
    s[c].font = Font(bold=True)

for ws in wb:
    for row in ws.iter_rows():
        for cell in row:
            cell.font = Font(name="Arial", bold=cell.font.bold, color=cell.font.color)

wb.save("stuck_parcels.xlsx")
print("saved stuck_parcels.xlsx")
